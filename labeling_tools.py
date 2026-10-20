"""Etiketleme araçları: klinisyen için Excel/CSV şablonu üretir, doldurulmuş dosyayı doğrulayıp
`cases_expert.json` dosyasına çevirir.

Kullanım
  python labeling_tools.py export --out etiket_sablonu.xlsx
  python labeling_tools.py import --in doldurulmus.xlsx --out cases_expert.json --labeler "Dr. Ad Soyad"

Şablon:
  * Mevcut vakalar (cases_legacy.py) hazır gelir; klinisyen yalnızca etiket sütunlarını doldurur.
  * K01..K30 satırları: risksiz KONTROL vakaları yazması için boş satırlar (FPR ölçümü için gerekir).
  * V01..V20 satırları: örtük / yazım hatalı / İngilizce / karışık dil / geçmiş zaman varyantları için boş satırlar.
  * "Açıklama" sayfasında alanların anlamı ve red_flags.py kategori anahtarları yer alır.
"""
import argparse
import difflib
import json
import os
import re
import sys
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

from cases import VALID_EMOTIONS, VALID_RISK, load_cases
from metrics import wilson_interval
from red_flags import RED_FLAG_RULES, _fold

COLUMNS = ["id", "user_message", "language", "emotion", "expected_red_flag", "expected_flags",
           "expected_risk", "adversarial_type", "variant_of", "labeler", "notes", "legacy_risk_level"]
EDITABLE = ["expected_red_flag", "expected_flags", "expected_risk", "adversarial_type", "variant_of", "labeler", "notes"]
ADVERSARIAL_TYPES = ["", "implicit", "typo", "code_switch", "english", "quoted", "past_tense", "negation",
                     "false_positive_trap", "colloquial", "prompt_injection", "control"]
LANGUAGES = ["tr", "en", "mixed", "de", "fr", "es", "ru"]
# Değerler _fold() sonrası (küçük harf, aksansız) karşılaştırılır
YES = {"evet", "e", "true", "1", "yes", "y", "x", "var"}
NO = {"hayir", "h", "false", "0", "no", "n", "yok"}
UNSURE = {"belirsiz", "?", "unclear", "unsure", "bilmiyorum"}


# ------------------------------------------------------------------
# Dışa aktarma
# ------------------------------------------------------------------
def _base_rows(include_starter: bool) -> List[Dict[str, Any]]:
    cases = load_cases(include_starter=include_starter, include_expert=False)
    rows = []
    for c in cases:
        legacy = re.search(r"legacy_risk_level=(\w+)", c.get("notes", ""))
        rows.append({
            "id": c["id"], "user_message": c["user_message"], "language": c.get("language", "tr"),
            "emotion": c.get("emotion", "Normal / Sakin"), "expected_red_flag": "", "expected_flags": "",
            "expected_risk": "", "adversarial_type": c.get("adversarial_type") or "", "variant_of": "",
            "labeler": "", "notes": "", "legacy_risk_level": legacy.group(1) if legacy else "",
        })
    return rows


def _blank_rows(prefix: str, n: int, emotion: str = "Normal / Sakin") -> List[Dict[str, Any]]:
    return [{"id": f"{prefix}{i:02d}", "user_message": "", "language": "tr", "emotion": emotion,
             "expected_red_flag": "", "expected_flags": "", "expected_risk": "", "adversarial_type": "",
             "variant_of": "", "labeler": "", "notes": "", "legacy_risk_level": ""} for i in range(1, n + 1)]


def template_dataframe(include_starter: bool = False, n_controls: int = 30, n_variants: int = 20) -> pd.DataFrame:
    rows = _base_rows(include_starter) + _blank_rows("K", n_controls) + _blank_rows("V", n_variants)
    return pd.DataFrame(rows, columns=COLUMNS)


def _explanation_rows() -> List[Tuple[str, str]]:
    return [
        ("id", "Vaka kimliği. Mevcut vakalarda DEĞİŞTİRMEYİN. K = yeni risksiz kontrol vakası, V = yeni varyant."),
        ("user_message", "Hasta mesajı. Mevcut vakaların metnini değiştirmeyin; K/V satırlarında yeni mesajı siz yazın."),
        ("language", "Mesajın dili: tr, en, mixed (karışık), de, fr, es, ru."),
        ("emotion", "Hastanın duygu durumu: " + " | ".join(VALID_EMOTIONS)),
        ("expected_red_flag", "EVET / HAYIR / BELİRSİZ. Bir klinisyen bu mesajda kırmızı bayrak düzeyinde (acil veya öncelikli) güvenlik yaklaşımı gerektiğini düşünüyor mu? "
                              "Risk içermeyen günlük sıkıntı mesajı = HAYIR. Emin değilseniz BELİRSİZ yazın (metriklere dahil edilmez)."),
        ("expected_flags", "Uygunsa kategori anahtarları, birden fazlaysa ';' ile ayırın (anahtarlar aşağıda). HAYIR ise boş bırakın."),
        ("expected_risk", "none / low / medium / high"),
        ("adversarial_type", "Zorluk türü: " + ", ".join(t for t in ADVERSARIAL_TYPES if t)),
        ("variant_of", "V satırlarında: varyantın türetildiği vakanın id'si (örn. 8)."),
        ("labeler", "Etiketleyen kişi (boşsa içe aktarırken varsayılan ad kullanılır)."),
        ("notes", "Serbest not (belirsizlik nedeni, klinik gerekçe vb.)."),
        ("legacy_risk_level", "Eski dosyadaki risk_level (yalnızca bilgi; içe aktarılmaz)."),
    ]


def export_template(out_path: str, include_starter: bool = False, n_controls: int = 30, n_variants: int = 20) -> str:
    """Şablonu .xlsx (açılır listeli) veya .csv olarak yazar."""
    df = template_dataframe(include_starter, n_controls, n_variants)
    if out_path.lower().endswith(".csv"):
        df.to_csv(out_path, index=False, encoding="utf-8-sig")
        return out_path

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Etiketleme"
    ws.append(COLUMNS)
    for rec in df.itertuples(index=False):
        ws.append(list(rec))

    head_fill = PatternFill("solid", fgColor="1F4E78")
    edit_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    widths = {"id": 7, "user_message": 70, "language": 9, "emotion": 15, "expected_red_flag": 14, "expected_flags": 28,
              "expected_risk": 12, "adversarial_type": 18, "variant_of": 10, "labeler": 16, "notes": 36, "legacy_risk_level": 12}
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths[col]
    last = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col in EDITABLE:
            row[COLUMNS.index(col)].fill = edit_fill
        # boş K/V satırlarında mesaj da düzenlenebilir
        if not str(row[1].value or ""):
            row[1].fill = edit_fill
    ws.freeze_panes = "C2"

    def add_list(col: str, items: List[str]) -> None:
        dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        letter = get_column_letter(COLUMNS.index(col) + 1)
        dv.add(f"{letter}2:{letter}{last}")

    add_list("expected_red_flag", ["EVET", "HAYIR", "BELİRSİZ"])
    add_list("expected_risk", list(VALID_RISK))
    add_list("emotion", list(VALID_EMOTIONS))
    add_list("language", LANGUAGES)
    add_list("adversarial_type", [t for t in ADVERSARIAL_TYPES if t])

    guide = wb.create_sheet("Açıklama")
    guide.append(["Alan", "Açıklama"])
    for k, v in _explanation_rows():
        guide.append([k, v])
    guide.append([])
    guide.append(["Kategori anahtarı (expected_flags)", "Açıklama / risk / sistem eylemi"])
    for key, rule in RED_FLAG_RULES.items():
        guide.append([key, f"{rule['description']} | risk: {rule['risk_level']} | eylem: {rule['action']}"])
    guide.column_dimensions["A"].width = 36
    guide.column_dimensions["B"].width = 120
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in guide[1]:
        cell.font = Font(bold=True)

    wb.save(out_path)
    return out_path


# ------------------------------------------------------------------
# İçe aktarma
# ------------------------------------------------------------------
def _read_table(path: str) -> pd.DataFrame:
    if path.lower().endswith((".xlsx", ".xlsm")):
        xl = pd.ExcelFile(path)
        sheet = "Etiketleme" if "Etiketleme" in xl.sheet_names else xl.sheet_names[0]
        df = xl.parse(sheet, dtype=str)
    else:
        df = pd.read_csv(path, sep=None, engine="python", encoding="utf-8-sig", dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    return df.fillna("")


def _parse_bool(value: str) -> Tuple[Optional[bool], Optional[str]]:
    v = _fold(value)  # Türkçe büyük/küçük harf güvenli ("BELİRSİZ", "HAYIR" ...)
    if v in YES:
        return True, None
    if v in NO:
        return False, None
    if v in UNSURE or v == "":
        return None, None
    return None, f"expected_red_flag anlaşılmadı: {value!r} (EVET/HAYIR/BELİRSİZ yazın)"


def _parse_flags(value: str) -> Tuple[List[str], List[str]]:
    flags, errors = [], []
    for token in re.split(r"[;,|\n]", value):
        key = token.strip().lower()
        if not key:
            continue
        if key in RED_FLAG_RULES:
            flags.append(key)
        else:
            close = difflib.get_close_matches(key, list(RED_FLAG_RULES), n=1)
            hint = f" (şunu mu demek istediniz: {close[0]}?)" if close else ""
            errors.append(f"bilinmeyen kategori {token.strip()!r}{hint}")
    return flags, errors


def import_labels(in_path: str, out_path: str = "cases_expert.json", default_labeler: str = "",
                  merge: bool = True) -> Dict[str, Any]:
    """Doldurulmuş şablonu doğrular ve `cases_expert.json` dosyasına yazar.

    Geçerli satırlar yazılır, hatalı satırlar `errors` içinde raporlanır (hiçbiri sessizce atlanmaz).
    BELİRSİZ veya boş etiketli satırlar dosyaya yazılmaz (metriklere girmez), sayısı raporlanır.
    """
    df = _read_table(in_path)
    missing_cols = [c for c in ("id", "user_message", "expected_red_flag") if c not in df.columns]
    if missing_cols:
        raise ValueError(f"Gerekli sütun(lar) yok: {missing_cols}")
    for col in COLUMNS:
        if col not in df.columns:
            df[col] = ""

    base = {c["id"]: c for c in load_cases(include_expert=False)}
    errors: List[Dict[str, Any]] = []
    warnings: List[Dict[str, Any]] = []
    cases: List[Dict[str, Any]] = []
    seen: set = set()
    n_unlabeled = n_empty = 0

    for idx, row in df.iterrows():
        line = idx + 2  # Excel satır numarası (başlık = 1)
        cid = str(row["id"]).strip()
        msg = str(row["user_message"]).strip()
        if not cid and not msg:
            n_empty += 1
            continue
        if not msg:
            n_empty += 1  # boş K/V satırı
            continue
        row_err: List[str] = []
        if not cid:
            row_err.append("id boş")
        if cid in seen:
            row_err.append(f"id tekrar ediyor: {cid}")
        seen.add(cid)

        flag, err = _parse_bool(str(row["expected_red_flag"]))
        if err:
            row_err.append(err)
        raw_label = _fold(str(row["expected_red_flag"]))
        flags, flag_errs = _parse_flags(str(row["expected_flags"]))
        row_err += flag_errs

        risk = str(row["expected_risk"]).strip().lower() or None
        if risk is not None and risk not in VALID_RISK:
            row_err.append(f"expected_risk geçersiz: {risk!r}")
        emotion = str(row["emotion"]).strip() or "Normal / Sakin"
        if emotion not in VALID_EMOTIONS:
            row_err.append(f"emotion geçersiz: {emotion!r}")
        language = str(row["language"]).strip().lower() or "tr"

        if row_err:
            errors.append({"satır": line, "id": cid, "hata": "; ".join(row_err)})
            continue

        if flag is None:
            if raw_label == "":
                is_new = cid[:1] in ("K", "V") and cid not in base
                if is_new:
                    errors.append({"satır": line, "id": cid, "hata": "yeni vaka yazılmış ama expected_red_flag boş"})
                else:
                    n_unlabeled += 1
            else:
                n_unlabeled += 1  # BELİRSİZ
            continue

        if cid in base and base[cid]["user_message"].strip() != msg:
            warnings.append({"satır": line, "id": cid, "uyarı": "mevcut vakanın mesajı değiştirilmiş"})
        if flag is False and flags:
            warnings.append({"satır": line, "id": cid, "uyarı": "HAYIR denmiş ama expected_flags dolu"})
        if flag is True and risk == "none":
            warnings.append({"satır": line, "id": cid, "uyarı": "EVET denmiş ama expected_risk=none"})
        if flag is False and risk == "high":
            warnings.append({"satır": line, "id": cid, "uyarı": "HAYIR denmiş ama expected_risk=high"})
        if cid.startswith("V") and cid not in base and not str(row["variant_of"]).strip():
            warnings.append({"satır": line, "id": cid, "uyarı": "varyant ama variant_of boş"})

        adv = str(row["adversarial_type"]).strip() or ("control" if (cid.startswith("K") and cid not in base) else None)
        note = str(row["notes"]).strip()
        if str(row["variant_of"]).strip():
            note = (note + f" variant_of={str(row['variant_of']).strip()}").strip()
        cases.append({
            "id": cid, "title": base.get(cid, {}).get("title", f"Uzman vakası {cid}"), "language": language,
            "user_message": msg, "emotion": emotion, "category": base.get(cid, {}).get("category", "expert_new"),
            "expected_red_flag": flag, "expected_flags": flags, "expected_risk": risk, "adversarial_type": adv,
            "label_status": "expert_labeled",
            "labeler": str(row["labeler"]).strip() or default_labeler, "source": os.path.basename(in_path), "notes": note,
        })

    merged: Dict[str, Dict[str, Any]] = {}
    if merge and os.path.exists(out_path):
        with open(out_path, "r", encoding="utf-8") as fh:
            existing = json.load(fh)
        for c in (existing.get("cases", []) if isinstance(existing, dict) else existing):
            merged[str(c["id"])] = c
    for c in cases:
        merged[c["id"]] = c

    payload = {
        "meta": {"imported_at": datetime.now().isoformat(timespec="seconds"), "source_file": os.path.basename(in_path),
                 "n_cases": len(merged)},
        "cases": list(merged.values()),
    }
    with open(out_path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)

    return {
        "out_path": out_path, "n_rows": len(df), "n_imported": len(cases), "n_unlabeled_or_unsure": n_unlabeled,
        "n_empty_rows": n_empty, "errors": errors, "warnings": warnings,
        "balance": label_balance(list(merged.values())),
    }


# ------------------------------------------------------------------
# Etiket dengesi / yeterlilik uyarıları
# ------------------------------------------------------------------
def label_balance(cases: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Etiketli vakaların dağılımı ve ölçümün ne kadar kesin olabileceğine dair uyarılar."""
    labeled = [c for c in cases if c.get("expected_red_flag") is not None]
    pos = [c for c in labeled if c["expected_red_flag"]]
    neg = [c for c in labeled if not c["expected_red_flag"]]
    by_lang: Dict[str, int] = {}
    by_adv: Dict[str, int] = {}
    for c in labeled:
        by_lang[c.get("language", "?")] = by_lang.get(c.get("language", "?"), 0) + 1
        a = c.get("adversarial_type") or "none"
        by_adv[a] = by_adv.get(a, 0) + 1

    notes: List[str] = []
    for name, n in (("pozitif", len(pos)), ("negatif (kontrol)", len(neg))):
        if n == 0:
            notes.append(f"{name} vaka yok: ilgili metrik hesaplanamaz.")
        else:
            # en kötü durum (oran 0.5) için yaklaşık %95 GA yarı genişliği
            _, lo, hi = wilson_interval(n // 2, n)
            half = (hi - lo) / 2
            if half > 0.10:
                notes.append(f"{name} n={n}: tahmini %95 GA yarı genişliği ≈ ±{half:.2f}; daha fazla vaka kesinliği artırır.")
    if len(by_lang) < 2:
        notes.append("Tek dilli veri: çok dilli sağlamlık (H3) test edilemez.")
    if not any(k != "none" for k in by_adv):
        notes.append("Zor vaka türü (örtük, yazım hatalı, İngilizce vb.) yok: kural motoru kırılganlığı ölçülemez.")
    return {"n_labeled": len(labeled), "n_positive": len(pos), "n_negative": len(neg),
            "by_language": by_lang, "by_adversarial_type": by_adv, "notes": notes}


def _print_report(rep: Dict[str, Any]) -> None:
    print(f"\nYazılan dosya: {rep['out_path']}")
    print(f"Satır: {rep['n_rows']} · içe aktarılan: {rep['n_imported']} · etiketsiz/belirsiz: {rep['n_unlabeled_or_unsure']} · boş: {rep['n_empty_rows']}")
    for e in rep["errors"]:
        print(f"  HATA   satır {e['satır']} ({e['id']}): {e['hata']}")
    for w in rep["warnings"]:
        print(f"  UYARI  satır {w['satır']} ({w['id']}): {w['uyarı']}")
    b = rep["balance"]
    print(f"Denge: pozitif={b['n_positive']} negatif={b['n_negative']} diller={b['by_language']} zorluk={b['by_adversarial_type']}")
    for n in b["notes"]:
        print("  •", n)


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="Vaka etiketleme şablonu / içe aktarma")
    sub = ap.add_subparsers(dest="cmd", required=True)
    ex = sub.add_parser("export", help="Şablon üret")
    ex.add_argument("--out", default="etiket_sablonu.xlsx")
    ex.add_argument("--include-starter", action="store_true", help="Sentetik başlangıç vakalarını da ekle")
    im = sub.add_parser("import", help="Doldurulmuş şablonu içe aktar")
    im.add_argument("--in", dest="inp", required=True)
    im.add_argument("--out", default="cases_expert.json")
    im.add_argument("--labeler", default="")
    im.add_argument("--no-merge", action="store_true")
    args = ap.parse_args(argv)

    if args.cmd == "export":
        print("Şablon yazıldı:", export_template(args.out, include_starter=args.include_starter))
        return 0
    rep = import_labels(args.inp, args.out, args.labeler, merge=not args.no_merge)
    _print_report(rep)
    return 1 if rep["errors"] else 0


if __name__ == "__main__":
    sys.exit(main())
